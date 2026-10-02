"""
Reading reviewer input back out of a report workbook.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook

ANNOTATIONS_SHEET = "Annotations"
GENERATED_SHEET = "_generated"
CANONICAL_COLUMNS = ("Canonical Target", "Canonical Situation")
REVIEWER_COLUMNS = ("Reviewer Verdict", "Reviewer Notes")


@dataclass
class ReviewRow:
    line_id: str
    text: str
    values: dict[str, Optional[str]] = field(default_factory=dict)
    generated: dict[str, Optional[str]] = field(default_factory=dict)

    def edited_canonical(self) -> dict[str, Optional[str]]:
        out = {}
        for col in CANONICAL_COLUMNS:
            if col not in self.values:
                continue
            value = _clean(self.values[col])
            if col not in self.generated:
                if value is not None:
                    out[col] = value
            elif value != _clean(self.generated[col]):
                out[col] = value
        return out


def _clean(value) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def read_review(path: Path) -> dict[str, ReviewRow]:
    if not path.exists():
        return {}
    wb = load_workbook(path, read_only=True, data_only=True)
    if ANNOTATIONS_SHEET not in wb.sheetnames:
        return {}
    rows = list(wb[ANNOTATIONS_SHEET].iter_rows(values_only=True))
    if not rows:
        return {}
    header = [str(h) if h is not None else "" for h in rows[0]]
    try:
        id_col, text_col = header.index("Line ID"), header.index("Humorous Text")
    except ValueError:
        return {}
    wanted = [c for c in (*CANONICAL_COLUMNS, *REVIEWER_COLUMNS) if c in header]

    review: dict[str, ReviewRow] = {}
    for row in rows[1:]:
        if not row or row[id_col] is None:
            continue
        line_id = str(row[id_col])
        review[line_id] = ReviewRow(
            line_id=line_id,
            text=str(row[text_col] or ""),
            values={c: row[header.index(c)] for c in wanted},
        )

    if GENERATED_SHEET in wb.sheetnames:
        gen_rows = list(wb[GENERATED_SHEET].iter_rows(values_only=True))
        if gen_rows:
            gen_header = [str(h) for h in gen_rows[0]]
            for row in gen_rows[1:]:
                record = dict(zip(gen_header, row))
                line_id = str(record.get("Line ID"))
                if line_id in review:
                    review[line_id].generated = {
                        c: record.get(c) for c in CANONICAL_COLUMNS if c in record
                    }
    wb.close()
    return review


def matching(review: dict[str, ReviewRow], line_id: str, text: str) -> Optional[ReviewRow]:
    row = review.get(line_id)
    if row is None or row.text.strip() != text.strip():
        return None
    return row
