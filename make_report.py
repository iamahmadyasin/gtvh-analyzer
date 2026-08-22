from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from schemas import Analysis

ROOT = Path(__file__).parent
INPUT_DIR = ROOT / "input"
OUTPUT_DIR = ROOT / "output"

FONT_NAME = "Arial"
HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF")
HUMOR_ROW_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
HYPERLINK_FONT = Font(name=FONT_NAME, color="0563C1", underline="single")
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")


def _style_header(ws, row: int, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = ws.cell(row=row + 1, column=1).coordinate


def _autofit(ws, widths: dict[int, int]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width


def build_report(analysis: Analysis, story_text: str, out_path: Path) -> None:
    n_lines = len(story_text.splitlines())

    # Map each global line number -> containing segment label
    def segment_for_line(line_no: int) -> str:
        for seg in analysis.segments:
            if seg.line_start <= line_no <= seg.line_end:
                return seg.label
        return ""

    # Map each global line number -> list of annotation line_ids touching it
    line_to_annotation_ids: dict[int, list[str]] = {i: [] for i in range(1, n_lines + 1)}
    for al in analysis.lines:
        for ln in range(al.span.line_start, al.span.line_end + 1):
            if ln in line_to_annotation_ids:
                line_to_annotation_ids[ln].append(al.line_id)

    annotation_row = {al.line_id: i + 2 for i, al in enumerate(analysis.lines)}

    wb = Workbook()

    # Story sheet
    story_ws = wb.active
    story_ws.title = "Story"
    story_headers = ["Line #", "Segment", "Text", "Annotation(s)"]
    for c, h in enumerate(story_headers, start=1):
        story_ws.cell(row=1, column=c, value=h)
    _style_header(story_ws, 1, len(story_headers))

    for i, raw_line in enumerate(story_text.splitlines(), start=1):
        row = i + 1
        story_ws.cell(row=row, column=1, value=i).alignment = TOP
        story_ws.cell(row=row, column=2, value=segment_for_line(i)).alignment = TOP
        text_cell = story_ws.cell(row=row, column=3, value=raw_line)
        text_cell.alignment = WRAP

        ids = line_to_annotation_ids.get(i, [])
        ann_cell = story_ws.cell(row=row, column=4)
        if ids:
            ann_cell.value = ", ".join(ids)
            # Jump to the first annotation touching this line
            first_row = annotation_row[ids[0]]
            ann_cell.hyperlink = f"#Annotations!A{first_row}"
            ann_cell.font = HYPERLINK_FONT
            for c in range(1, len(story_headers) + 1):
                story_ws.cell(row=row, column=c).fill = HUMOR_ROW_FILL

    _autofit(story_ws, {1: 8, 2: 24, 3: 90, 4: 18})
    story_ws.auto_filter.ref = f"A1:D{n_lines + 1}"

    # Annotations sheet
    ann_ws = wb.create_sheet("Annotations")
    ann_headers = [
        "Line ID", "Context ▶", "Classification", "Narrative Level",
        "Segment", "Line Type", "Span", "Humorous Text", "Disjunctor",
        "Script 1", "Script 2", "Binary Category", "Opposition Type",
        "Situation", "Target", "Orientation", "Narrative Strategy",
        "Wordplay?", "Wordplay Level", "Wordplay Subtype", "Wordplay Note",
        "Register Effect?", "Register Subtype", "Register Note",
        "Justification", "Reviewer Verdict", "Reviewer Notes",
    ]
    for c, h in enumerate(ann_headers, start=1):
        ann_ws.cell(row=1, column=c, value=h)
    _style_header(ann_ws, 1, len(ann_headers))

    seg_label = {s.segment_id: s.label for s in analysis.segments}

    for i, al in enumerate(analysis.lines):
        row = i + 2
        a = al.annotation
        so = a.script_opposition
        lang = a.language
        values = [
            al.line_id,
            None,
            a.classification.value,
            a.narrative_level_of_classification.value,
            seg_label.get(al.segment_id, al.segment_id),
            al.line_type.value,
            f"{al.span.line_start}\u2013{al.span.line_end}",
            al.span.text,
            al.disjunctor or "",
            so.script_1,
            so.script_2,
            so.essential_binary_category.value,
            so.opposition_type.value,
            a.situation,
            a.target or "",
            a.orientation.value,
            a.narrative_strategy,
            "Yes" if lang.is_wordplay else "No",
            lang.wordplay_level.value if lang.wordplay_level else "",
            lang.wordplay_subtype or "",
            lang.wordplay_note,
            "Yes" if lang.is_register_effect else "No",
            lang.register_effect_subtype or "",
            lang.register_note,
            a.justification,
            "",  # Reviewer Verdict — blank for user
            "",  # Reviewer Notes — blank for user
        ]
        for c, v in enumerate(values, start=1):
            cell = ann_ws.cell(row=row, column=c, value=v)
            cell.alignment = WRAP if c in (8, 25, 27) else TOP

        # Hyperlink back to the first line of the story span
        ctx_cell = ann_ws.cell(row=row, column=2, value="\u25b6 view in story")
        ctx_cell.hyperlink = f"#Story!A{al.span.line_start + 1}"
        ctx_cell.font = HYPERLINK_FONT

    _autofit(ann_ws, {
        1: 10, 2: 14, 3: 12, 4: 14, 5: 22, 6: 14, 7: 10, 8: 45, 9: 20,
        10: 22, 11: 22, 12: 18, 13: 20, 14: 20, 15: 20, 16: 12, 17: 18,
        18: 10, 19: 14, 20: 18, 21: 30, 22: 14, 23: 18, 24: 30, 25: 40,
        26: 16, 27: 30,
    })
    ann_ws.auto_filter.ref = f"A1:{get_column_letter(len(ann_headers))}{len(analysis.lines) + 1}"

    # Reviewer Verdict dropdown
    dv = DataValidation(type="list", formula1='"Agree,Disagree,Partial,Unsure"', allow_blank=True)
    ann_ws.add_data_validation(dv)
    dv.add(f"Z2:Z{len(analysis.lines) + 1}")

    # Segments sheet
    seg_ws = wb.create_sheet("Segments")
    seg_headers = ["Segment ID", "Label", "Narrative Level", "Line Range",
                   "Parent", "Terminal?", "Cue", "Description"]
    for c, h in enumerate(seg_headers, start=1):
        seg_ws.cell(row=1, column=c, value=h)
    _style_header(seg_ws, 1, len(seg_headers))
    for i, seg in enumerate(analysis.segments):
        row = i + 2
        values = [
            seg.segment_id, seg.label, seg.narrative_level.value,
            f"{seg.line_start}\u2013{seg.line_end}",
            seg.parent_segment_id or "", "Yes" if seg.is_terminal else "No",
            seg.segmentation_cue, seg.description,
        ]
        for c, v in enumerate(values, start=1):
            seg_ws.cell(row=row, column=c, value=v).alignment = TOP
    _autofit(seg_ws, {1: 10, 2: 26, 3: 16, 4: 12, 5: 12, 6: 10, 7: 22, 8: 40})

    # Apply the base font everywhere
    for ws in (story_ws, ann_ws, seg_ws):
        for row in ws.iter_rows():
            for cell in row:
                if cell.font is None or cell.font.name != FONT_NAME:
                    existing = cell.font
                    cell.font = Font(
                        name=FONT_NAME,
                        bold=existing.bold, color=existing.color,
                        underline=existing.underline,
                    )

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def _run_one(json_path: Path, story_path: Path | None, out_path: Path | None) -> None:
    data = json.loads(json_path.read_text(encoding="utf-8"))
    analysis = Analysis.model_validate(data)

    if story_path is None:
        story_path = INPUT_DIR / analysis.source_filename
    if not story_path.exists():
        print(f"  \u2717 story text not found: {story_path}", file=sys.stderr)
        print(f"    (expected from source_filename={analysis.source_filename!r} "
              f"in {json_path.name})", file=sys.stderr)
        return

    if out_path is None:
        out_path = json_path.with_suffix(".xlsx")

    story_text = story_path.read_text(encoding="utf-8")
    build_report(analysis, story_text, out_path)
    print(f"  wrote {out_path}")


def main() -> None:
    parser = argparse.ArgumentParser(description="Build a reviewable Excel report.")
    parser.add_argument("--json", type=Path, help="Single analysis JSON (default: all in output/).")
    parser.add_argument("--story", type=Path, help="Override path to the original story .txt.")
    parser.add_argument("--out", type=Path, help="Output .xlsx path (default: alongside the JSON).")
    args = parser.parse_args()

    if args.json:
        print(f"\u2192 Building report for {args.json.name}")
        _run_one(args.json, args.story, args.out)
        return

    files = sorted(OUTPUT_DIR.glob("*.json"))
    if not files:
        print(f"No .json files in {OUTPUT_DIR}. Run the pipeline first.")
        return
    for jp in files:
        print(f"\u2192 Building report for {jp.name}")
        _run_one(jp, None, None)


if __name__ == "__main__":
    main()
