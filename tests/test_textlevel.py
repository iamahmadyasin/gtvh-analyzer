"""Tests for the deterministic parts of the text-level stage.

Run with:  python -m unittest discover tests
No API key or network needed.
"""

import asyncio
import tempfile
import unittest
from pathlib import Path

from normalize import NormalizationParams, group_labels, normalize_analysis, string_similarity
from review_workbook import read_review
from schemas import (
    Analysis, AnnotatedLine, KRAnnotation, LanguageKR, NarrativeSegment,
    ScriptOpposition, TargetEntry, TextLevelReport, TextSpan,
)
from textlevel import TextLevelParams, compute_metrics

WORDS_PER_LINE = 50


def _story(n_lines: int) -> str:
    return "\n".join(" ".join(f"w{i}x{j}" for j in range(WORDS_PER_LINE)) for i in range(n_lines))


def _line(n: int, para: int, target=None, target_id=None, situation="cotext",
          classification="jab", opposition="normal_vs_abnormal", text=None) -> AnnotatedLine:
    return AnnotatedLine(
        line_id=f"HL-{n:03d}",
        span=TextSpan(line_start=para, line_end=para, text=text or f"w{para - 1}x10 w{para - 1}x11"),
        segment_id="NS-01", line_type="discrete", confidence="high",
        annotation=KRAnnotation(
            reasoning="", line_id=f"HL-{n:03d}", classification=classification,
            narrative_level_of_classification="level_0",
            script_opposition=ScriptOpposition(
                script_1="A", script_2="B", essential_binary_category="none",
                opposition_type=opposition),
            situation=situation, orientation="other", target_id=target_id, target=target,
            narrative_strategy="narration", narrative_strategy_note=None,
            language=LanguageKR(is_wordplay=False, wordplay_level=None, wordplay_subtype=None,
                                is_register_effect=False, register_effect_subtype=None)))


def _analysis(lines, n_paras=100, inventory=()) -> Analysis:
    seg = NarrativeSegment(segment_id="NS-01", label="All", narrative_level="level_0",
                           line_start=1, line_end=n_paras, parent_segment_id=None,
                           is_terminal=False, segmentation_cue="", description="")
    return Analysis(source_filename="t.txt", segments=[seg], lines=list(lines),
                    target_inventory=list(inventory))


class NormalizationTests(unittest.TestCase):
    def test_sentinels_never_merge_and_paraphrases_do(self):
        a = _analysis([
            _line(1, 1, situation="Lady B's party"), _line(2, 2, situation="Lady B's party (?)"),
            _line(3, 3, situation="party at Lady B's"), _line(4, 4, situation="cotext"),
            _line(5, 5, situation="irr"), _line(6, 6, situation="Cotext"),
        ])
        asyncio.run(normalize_analysis(a, NormalizationParams(method="string")))
        got = [ln.canonical_situation for ln in a.lines]
        self.assertEqual(got, ["Lady B's party"] * 3 + ["cotext", "irr", "cotext"])

    def test_new_target_maps_to_inventory_alias(self):
        entry = TargetEntry(target_id="T-01", label="Lord Arthur Savile", aliases=["Lord Arthur"],
                            kind="person", social_class="upper", sphere="high_society",
                            description="")
        a = _analysis([_line(1, 1, target="lord arthur", target_id=None),
                       _line(2, 2, target="the vicar", target_id=None),
                       _line(3, 3, target="The Vicar (?)", target_id=None)], inventory=[entry])
        asyncio.run(normalize_analysis(a, NormalizationParams(method="string")))
        self.assertEqual((a.lines[0].canonical_target, a.lines[0].canonical_target_id),
                         ("Lord Arthur Savile", "T-01"))
        self.assertEqual(a.lines[1].canonical_target, a.lines[2].canonical_target)
        self.assertIsNone(a.lines[1].canonical_target_id)

    def test_grouping_does_not_chain(self):
        from collections import Counter
        mapping = group_labels(Counter({"dinner party": 3, "garden party": 2}),
                               string_similarity, 0.8)
        self.assertNotEqual(mapping["dinner party"], mapping["garden party"])


class DistributionTests(unittest.TestCase):
    def _metrics(self, paras, **params):
        lines = [_line(i + 1, p) for i, p in enumerate(paras)]
        return compute_metrics(_analysis(lines), _story(100),
                               TextLevelParams(n_simulations=2000, **params))

    def test_sections_and_ratios(self):
        m = self._metrics([1, 2, 3, 51], n_sections=4)
        d = m.distribution
        self.assertEqual(d.n_words, 5000)
        self.assertEqual([s.n_lines for s in d.sections], [3, 0, 1, 0])
        self.assertEqual(d.words_per_line, 1250.0)
        self.assertEqual(d.sections[0].words_per_line, round(1250 / 3, 1))
        self.assertIsNone(d.sections[1].words_per_line)

    def test_clustered_rejects_random(self):
        m = self._metrics([1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 90, 91, 92, 93, 94, 95])
        random_test = next(t for t in m.distribution.tests if t.null_hypothesis == "random")
        self.assertLess(random_test.p_value_greater, 0.05)
        self.assertTrue(random_test.conclusion.startswith("reject: lines are more clustered"))
        self.assertTrue(m.distribution.waves)
        self.assertTrue(m.distribution.serious_reliefs)

    def test_even_spacing_reads_as_more_regular_than_random(self):
        m = self._metrics(list(range(3, 100, 5)))
        random_test = next(t for t in m.distribution.tests if t.null_hypothesis == "random")
        self.assertLess(random_test.p_value_less, 0.05)
        uniform = next(t for t in m.distribution.tests if t.null_hypothesis == "uniform")
        self.assertTrue(uniform.conclusion.startswith("not rejected"))
        self.assertFalse(m.distribution.serious_reliefs)

    def test_reproducible_with_seed(self):
        a = self._metrics([1, 5, 9, 40, 41, 80]).distribution.tests
        b = self._metrics([1, 5, 9, 40, 41, 80]).distribution.tests
        self.assertEqual(a, b)


class StrandTests(unittest.TestCase):
    def test_central_peripheral_combs_bridges(self):
        lines = (
            [_line(i + 1, p, target="Hero") for i, p in enumerate([2, 30, 60, 95])]
            + [_line(10 + i, p, target="Hostess") for i, p in enumerate([40, 41, 42])]
        )
        m = compute_metrics(_analysis(lines), _story(100),
                            TextLevelParams(n_simulations=200, strand_pairs="none"))
        by_key = {s.key: s for s in m.strands}
        hero, hostess = by_key["target=Hero"], by_key["target=Hostess"]
        self.assertEqual(hero.centrality, "central")
        self.assertEqual(hostess.centrality, "peripheral")
        self.assertEqual(len(hostess.comb_ids), 1)
        self.assertEqual(hero.comb_ids, [])
        self.assertEqual(len(hero.bridge_ids), 3)   # 2->30, 30->60, 60->95 all >= 0.25 apart
        self.assertAlmostEqual(hero.share, 4 / 7, places=3)

    def test_min_lines_and_excluded_values(self):
        lines = [_line(i + 1, p, situation="cotext") for i, p in enumerate([1, 2])]
        m = compute_metrics(_analysis(lines), _story(100), TextLevelParams(n_simulations=100))
        self.assertFalse(any(s.key.startswith("situation=") for s in m.strands))
        self.assertFalse(any(s.n_lines < 3 for s in m.strands))

    def test_identical_line_sets_merge(self):
        lines = [_line(i + 1, p, target="Hero", opposition="possible_vs_impossible")
                 for i, p in enumerate([5, 50, 90])]
        m = compute_metrics(_analysis(lines), _story(100), TextLevelParams(n_simulations=100))
        # target, orientation, opposition type and narrative strategy all pick
        # out the same three lines: one strand, named by its target.
        self.assertEqual(len(m.strands), 1)
        self.assertEqual(m.strands[0].key, "target=Hero")
        self.assertIn("opposition_type=possible_vs_impossible", m.strands[0].equivalent_keys)

    def test_output_round_trips(self):
        lines = [_line(i + 1, p, target="Hero") for i, p in enumerate([5, 50, 90])]
        m = compute_metrics(_analysis(lines), _story(100), TextLevelParams(n_simulations=100))
        again = TextLevelReport.model_validate_json(TextLevelReport(metrics=m).model_dump_json())
        self.assertEqual(again.metrics, m)
        self.assertEqual(m.params["min_strand_lines"], 3)


class LegacyAnalysisTests(unittest.TestCase):
    def test_unnormalized_analysis_inside_event_loop(self):
        # analyze_text.py calls compute_metrics from inside asyncio.run();
        # analyses from before normalization must still work there.
        lines = [_line(i + 1, p, target="Hero", situation="ball (?)") for i, p in enumerate([5, 50, 90])]
        a = _analysis(lines)
        self.assertIsNone(a.lines[0].canonical_situation)

        async def run():
            return compute_metrics(a, _story(100), TextLevelParams(n_simulations=100))

        m = asyncio.run(run())
        self.assertEqual(m.lines[0].features["situation"], "ball")
        self.assertIn("string similarity", m.params["normalization"])


class ReviewRoundTripTests(unittest.TestCase):
    def test_edits_survive_rebuild_and_feed_text_level(self):
        import make_report
        from openpyxl import load_workbook
        lines = [_line(i + 1, p, target="Hero", situation="ball") for i, p in enumerate([5, 50, 90])]
        a = _analysis(lines)
        asyncio.run(normalize_analysis(a, NormalizationParams(method="string")))
        with tempfile.TemporaryDirectory() as tmp:
            xlsx = Path(tmp) / "t.xlsx"
            make_report.build_report(a, _story(100), xlsx)
            wb = load_workbook(xlsx)
            ws = wb["Annotations"]
            hdr = [c.value for c in ws[1]]
            ws.cell(2, hdr.index("Canonical Target") + 1).value = "Villain"
            ws.cell(3, hdr.index("Canonical Situation") + 1).value = None   # cleared
            wb.save(xlsx)

            rows = read_review(xlsx)
            self.assertEqual(rows["HL-001"].edited_canonical(), {"Canonical Target": "Villain"})
            self.assertEqual(rows["HL-002"].edited_canonical(), {"Canonical Situation": None})
            self.assertEqual(rows["HL-003"].edited_canonical(), {})

            kept = make_report.build_report(a, _story(100), xlsx)
            self.assertEqual(kept, 2)
            m = compute_metrics(a, _story(100), TextLevelParams(n_simulations=100),
                                review_path=xlsx)
            feats = {lf.line_id: lf.features for lf in m.lines}
            self.assertEqual(feats["HL-001"]["target"], "Villain")
            self.assertIsNone(feats["HL-002"]["situation"])
            self.assertEqual(m.reviewer_overrides, 2)

    def test_stale_rows_are_ignored(self):
        import make_report
        a = _analysis([_line(1, 5, target="Hero")])
        asyncio.run(normalize_analysis(a, NormalizationParams(method="string")))
        with tempfile.TemporaryDirectory() as tmp:
            xlsx = Path(tmp) / "t.xlsx"
            make_report.build_report(a, _story(100), xlsx)
            from openpyxl import load_workbook
            wb = load_workbook(xlsx)
            ws = wb["Annotations"]
            hdr = [c.value for c in ws[1]]
            ws.cell(2, hdr.index("Canonical Target") + 1).value = "Villain"
            wb.save(xlsx)
            b = _analysis([_line(1, 7, target="Hero", text="a different line")])
            asyncio.run(normalize_analysis(b, NormalizationParams(method="string")))
            m = compute_metrics(b, _story(100), TextLevelParams(n_simulations=100), review_path=xlsx)
            self.assertEqual(m.lines[0].features["target"], "Hero")
            self.assertEqual(m.reviewer_overrides, 0)


if __name__ == "__main__":
    unittest.main()
